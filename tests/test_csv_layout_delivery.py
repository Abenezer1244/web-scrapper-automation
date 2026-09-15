"""Versioned CSV layout, wired end to end (scraper config -> download / scheduled export).

Real DB + real endpoints (conftest `db` / `client` / token fixtures), no mocks. Pins
the compatibility contract the owner chose: existing scrapers keep the legacy_v1
headers, new scrapers get crm_v1, and no edit silently switches a config's layout.
"""
import csv
import io
import uuid
from datetime import datetime

import pytest
from openpyxl import load_workbook

import src.db.session as _db_session
from src.api.routes.scrapers import _merge_deliver
from src.api.schemas import DeliverConfig, DeliverUpdate, ScraperConfigResponse
from src.db.models import Job, Result, ScraperConfig
from src.utils.data_exporter import DataExporter
from src.utils.lead_export import (
    CRM_V1_LABELS,
    EXPORT_LAYOUTS,
    LAYOUT_CRM_V1,
    LAYOUT_LEGACY_V1,
    LEAD_CSV_COLUMNS,
    resolve_export_layout,
    resolve_lead_export_columns,
)
from src.workers.tasks import _result_rows_to_export_dicts


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


# ─── Schema + merge (no DB) ───────────────────────────────────────────────────

def test_schema_literals_match_exporter_layouts():
    assert set(DeliverConfig.model_fields["csv_layout"].annotation.__args__[0].__args__) == set(
        EXPORT_LAYOUTS
    )


def test_invalid_layout_rejected():
    with pytest.raises(ValueError):
        DeliverConfig(csv_layout="crm_v9")


@pytest.mark.parametrize("stored, expected", [
    ({}, None),                              # pre-existing config: stays unset (legacy)
    ({"csv_layout": "legacy_v1"}, "legacy_v1"),
    ({"csv_layout": "crm_v1"}, "crm_v1"),
    ({"csv_layout": "garbage"}, None),       # never re-validate a bad stored value into a 422
])
def test_edit_without_layout_keeps_stored(stored, expected):
    merged = _merge_deliver(stored, DeliverUpdate(formats=["csv"]))
    assert merged.csv_layout == expected


def test_edit_with_layout_switches():
    merged = _merge_deliver({"csv_layout": "legacy_v1"}, DeliverUpdate(csv_layout="crm_v1"))
    assert merged.csv_layout == "crm_v1"


def test_edit_form_echo_of_get_response_is_accepted():
    # GET returns the effective layout; DeliverUpdate is extra="forbid", so a verbatim
    # echo must validate or every scraper edit would 422.
    DeliverUpdate(formats=["csv"], csv_layout="legacy_v1", webhook_secret_set=False)


@pytest.mark.parametrize("stored, effective", [
    ({}, "legacy_v1"), ({"csv_layout": None}, "legacy_v1"),
    ({"csv_layout": "crm_v1"}, "crm_v1"), ({"csv_layout": "bogus"}, "legacy_v1"),
])
def test_response_reports_effective_layout(stored, effective):
    now = datetime.now()
    resp = ScraperConfigResponse(
        id="x", user_id="u", name="n", county="king", state="WA", record_type="probate",
        fields={}, enrichment={}, schedule={}, deliver=stored, active=True,
        created_at=now, updated_at=now,
    )
    assert resp.deliver["csv_layout"] == effective
    columns, labels = resolve_export_layout(stored.get("csv_layout"), "probate")
    assert (labels is not None) == (effective == LAYOUT_CRM_V1)


# ─── Scheduled export projection ──────────────────────────────────────────────

def test_scheduled_projection_carries_stored_situs():
    row = Result(
        id=str(uuid.uuid4()), job_id=str(uuid.uuid4()), user_id=str(uuid.uuid4()),
        party_name="SMITH JOHN", property_address="123 MAIN ST",
        property_city="SEATTLE", property_state="WA", property_zip="98101",
    )
    (projected,) = _result_rows_to_export_dicts([row])
    assert (projected["property_city"], projected["property_state"], projected["property_zip"]) == (
        "SEATTLE", "WA", "98101")


def test_scheduled_crm_csv_matches_values(tmp_path):
    row = Result(
        id=str(uuid.uuid4()), job_id=str(uuid.uuid4()), user_id=str(uuid.uuid4()),
        party_name="SHIRLEY A JOHNSON", property_address="24910 51ST AVE E",
        property_city="GRAHAM", property_state="WA", property_zip="98338",
        mailing_address="PO BOX 465, LAKEBAY, WA, 98349-0465", parcel_id="0007200015",
        phone="2535551111", phones=[{"number": "2535551111", "type": "Mobile"}],
    )
    columns, labels = resolve_export_layout(LAYOUT_CRM_V1, "trustee_sale")
    path = DataExporter(export_dir=str(tmp_path)).export(
        _result_rows_to_export_dicts([row]), fmt="csv", columns=columns, labels=labels,
        context={"county": "pierce", "state": "WA", "record_type": "trustee_sale"},
    )
    with open(path, newline="", encoding="utf-8") as f:
        (lead,) = list(csv.DictReader(f))
    assert (lead["First Name"], lead["Last Name"]) == ("SHIRLEY", "JOHNSON")
    assert (lead["Property Address"], lead["Property City"], lead["Property Zip"]) == (
        "24910 51ST AVE E", "GRAHAM", "98338")
    assert (lead["Mailing Address"], lead["Mailing City"], lead["Mailing Zip"]) == (
        "PO BOX 465", "LAKEBAY", "98349-0465")
    assert lead["Parcel ID"] == "0007200015" and lead["Phone 1"] == "2535551111"
    assert (lead["County"], lead["Record Type"]) == ("Pierce", "Trustee Sale")


def test_excel_uses_labels_and_text_identifier_cells(tmp_path):
    rec = {"party_name": "SMITH JOHN", "parcel_id": "0040000055",
           "property_address": "1 ELM ST, HOLTSVILLE, NY 00501"}
    columns, labels = resolve_export_layout(LAYOUT_CRM_V1, "probate")
    path = DataExporter(export_dir=str(tmp_path)).export(
        [rec], fmt="excel", columns=columns, labels=labels,
        context={"county": "king", "state": "WA", "record_type": "probate"},
    )
    ws = load_workbook(path)["Leads"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["First Name", "Last Name", "Party Name"]
    parcel = ws.cell(row=2, column=header.index("Parcel ID") + 1)
    zip_cell = ws.cell(row=2, column=header.index("Property Zip") + 1)
    assert parcel.value == "0040000055" and parcel.data_type == "s" and parcel.number_format == "@"
    assert zip_cell.value == "00501" and zip_cell.number_format == "@"


def test_json_keeps_snake_case_keys_with_context(tmp_path):
    import json

    columns, _labels = resolve_export_layout(LAYOUT_CRM_V1, "probate")
    path = DataExporter(export_dir=str(tmp_path)).export(
        [{"party_name": "SMITH JOHN"}], fmt="json", columns=columns,
        context={"county": "king", "state": "WA", "record_type": "probate"},
    )
    (row,) = json.loads(open(path, encoding="utf-8").read())
    assert list(row)[:3] == ["first_name", "last_name", "party_name"]
    assert row["first_name"] == "JOHN" and row["record_type"] == "Probate"


# ─── Endpoints (real DB) ──────────────────────────────────────────────────────

async def _config(user_id: str, **overrides) -> ScraperConfig:
    data = {
        "id": str(uuid.uuid4()), "user_id": user_id, "name": "Layout Test",
        "county": "pierce", "state": "WA", "record_type": "pre_foreclosure",
        "fields": {"party_name": True}, "enrichment": {}, "schedule": {"frequency": "manual"},
        "deliver": {"formats": ["csv"], "emails": []},
    }
    data.update(overrides)
    async with _db_session.AsyncSessionLocal() as s:
        cfg = ScraperConfig(**data)
        s.add(cfg)
        await s.commit()
        await s.refresh(cfg)
        return cfg


async def _done_job_with_lead(user_id: str, config: ScraperConfig, **lead) -> str:
    job_id = str(uuid.uuid4())
    async with _db_session.AsyncSessionLocal() as s:
        s.add(Job(id=job_id, user_id=user_id, scraper_config_id=config.id, status="done",
                  trigger="manual", export_key=f"exports/{job_id}.csv"))
        await s.flush()
        s.add(Result(id=str(uuid.uuid4()), job_id=job_id, user_id=user_id, **lead))
        await s.commit()
    return job_id


async def _download(client, job_id: str, token: str):
    url = await client.get(f"/jobs/{job_id}/export-url", headers=_auth(token))
    assert url.status_code == 200, url.text
    return await client.get(url.json()["url"])


_LEAD = {
    "party_name": "HALL MARVIN WAYNE / HALL JUNE", "property_address": "510 106TH ST S",
    "property_city": "TACOMA", "property_state": "WA", "property_zip": "98444",
    "mailing_address": "510 106TH ST S, TACOMA, WA, 98444", "parcel_id": "0007200015",
}


async def test_download_crm_layout_headers_and_values(client, db, starter_user, starter_token):
    cfg = await _config(starter_user.id, deliver={"formats": ["csv"], "csv_layout": "crm_v1"})
    job_id = await _done_job_with_lead(starter_user.id, cfg, **_LEAD)
    resp = await _download(client, job_id, starter_token)
    assert resp.status_code == 200
    rows = list(csv.reader(io.StringIO(resp.text)))
    columns, labels = resolve_export_layout(LAYOUT_CRM_V1, "pre_foreclosure")
    assert rows[0] == [labels[c] for c in columns]
    lead = dict(zip(rows[0], rows[1], strict=True))
    assert len(rows) == 2  # one lead, one row
    assert (lead["First Name"], lead["Last Name"], lead["Party Name"]) == (
        "MARVIN", "HALL", "HALL MARVIN WAYNE / HALL JUNE")
    assert (lead["Property Address"], lead["Property City"], lead["Property State"],
            lead["Property Zip"]) == ("510 106TH ST S", "TACOMA", "WA", "98444")
    assert (lead["Mailing Address"], lead["Mailing City"], lead["Mailing Zip"]) == (
        "510 106TH ST S", "TACOMA", "98444")
    assert lead["Parcel ID"] == "0007200015"
    assert (lead["County"], lead["County State"], lead["Record Type"]) == (
        "Pierce", "WA", "Pre-Foreclosure")


async def test_download_existing_config_keeps_legacy_headers(client, db, starter_user, starter_token):
    cfg = await _config(starter_user.id)  # stored deliver has no csv_layout
    job_id = await _done_job_with_lead(starter_user.id, cfg, **_LEAD)
    resp = await _download(client, job_id, starter_token)
    rows = list(csv.reader(io.StringIO(resp.text)))
    assert rows[0] == resolve_lead_export_columns("pre_foreclosure")
    assert set(rows[0]) <= set(LEAD_CSV_COLUMNS)
    lead = dict(zip(rows[0], rows[1], strict=True))
    # Legacy file gets the corrected, source-aware name split through the config.
    assert (lead["first_name"], lead["last_name"]) == ("MARVIN", "HALL")
    assert not set(CRM_V1_LABELS.values()) & set(rows[0])


async def test_download_is_tenant_isolated(client, db, starter_user, business_user, starter_token):
    cfg = await _config(business_user.id, deliver={"formats": ["csv"], "csv_layout": "crm_v1"})
    job_id = await _done_job_with_lead(business_user.id, cfg, **_LEAD)
    url = await client.get(f"/jobs/{job_id}/export-url", headers=_auth(starter_token))
    assert url.status_code == 404
    direct = await client.get(f"/jobs/{job_id}/download", headers=_auth(starter_token))
    assert direct.status_code == 404
    assert "HALL MARVIN" not in direct.text


async def test_create_stamps_crm_and_edit_keeps_it(client, db, starter_user, starter_token):
    created = await client.post("/scrapers", json={
        "name": "New Layout", "county": "pierce", "state": "WA", "record_type": "probate",
    }, headers=_auth(starter_token))
    assert created.status_code == 201, created.text
    assert created.json()["deliver"]["csv_layout"] == LAYOUT_CRM_V1
    body = created.json()
    # Edit that omits csv_layout (a client that predates the field) must not flip it.
    edited = await client.patch(f"/scrapers/{body['id']}", json={
        "updated_at": body["updated_at"], "deliver": {"formats": ["csv"], "emails": []},
    }, headers=_auth(starter_token))
    assert edited.status_code == 200, edited.text
    assert edited.json()["deliver"]["csv_layout"] == LAYOUT_CRM_V1


async def test_edit_of_existing_config_stays_legacy_and_echo_is_accepted(
    client, db, starter_user, starter_token
):
    cfg = await _config(starter_user.id)
    got = await client.get(f"/scrapers/{cfg.id}", headers=_auth(starter_token))
    assert got.status_code == 200, got.text
    assert got.json()["deliver"]["csv_layout"] == LAYOUT_LEGACY_V1
    echoed = await client.patch(f"/scrapers/{cfg.id}", json={
        "updated_at": got.json()["updated_at"], "name": "Renamed",
        "deliver": got.json()["deliver"],
    }, headers=_auth(starter_token))
    assert echoed.status_code == 200, echoed.text
    assert echoed.json()["deliver"]["csv_layout"] == LAYOUT_LEGACY_V1
